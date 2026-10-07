import io
import re
import secrets
import string
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Query,
    Response,
    UploadFile,
    status
)
from sqlalchemy import or_, text
from sqlalchemy.orm import Session

from app.core.email import send_new_account_email
from app.core.permissions import require_permission
from app.core.security import hash_password
from app.database import get_db
from app.models.password_reset import PasswordResetToken
from app.models.role import Role
from app.models.session import UserSession
from app.models.user import User
from app.models.user_role import UserRole
from app.schemas.user import (
    ConfirmImportRequest,
    CreateUserRequest,
    ImportConfirmUserItem,
    ImportPreviewResponse,
    ImportSummaryResponse,
    ImportUserRow,
    LockUserRequest,
    UpdateUserRequest
)


router = APIRouter(
    prefix="/api/users",
    tags=["User Management"]
)


def generate_temporary_password() -> str:
    alphabet = (
        string.ascii_letters
        + string.digits
    )

    return "".join(
        secrets.choice(alphabet)
        for _ in range(12)
    )


@router.post("")
async def create_user(
    data: CreateUserRequest,
    current_user: User = Depends(
        require_permission("USER_MANAGE")
    ),
    db: Session = Depends(get_db)
):
    existing_email = db.query(User).filter(
        User.email == data.email
    ).first()

    if existing_email:
        raise HTTPException(
            status_code=400,
            detail="Email đã tồn tại trong hệ thống"
        )

    existing_username = db.query(User).filter(
        User.username == data.username
    ).first()

    if existing_username:
        raise HTTPException(
            status_code=400,
            detail="Tên đăng nhập đã tồn tại"
        )

    role = db.query(Role).filter(
        Role.role_id == data.role_id
    ).first()

    if role is None:
        raise HTTPException(
            status_code=404,
            detail="Vai trò không tồn tại"
        )

    temporary_password = (
        generate_temporary_password()
    )

    user = User(
        username=data.username,
        password=hash_password(
            temporary_password
        ),
        full_name=data.full_name,
        email=data.email,
        phone=data.phone,
        role_id=data.role_id,
        failed_login_attempts=0,
        is_locked=False,
        lock_reason=None
    )

    db.add(user)
    db.flush()

    user_role = UserRole(
        user_id=user.user_id,
        role_id=data.role_id
    )

    db.add(user_role)
    db.commit()
    db.refresh(user)

    await send_new_account_email(
        email=user.email,
        full_name=user.full_name,
        temporary_password=temporary_password
    )

    return {
        "message": "Tạo tài khoản thành công",
        "user_id": user.user_id,
        "email": user.email,
        "role": role.role_name
    }


@router.put("/{user_id}")
def update_user(
    user_id: int,
    data: UpdateUserRequest,
    current_user: User = Depends(
        require_permission("USER_MANAGE")
    ),
    db: Session = Depends(get_db)
):
    if current_user.user_id == user_id:
        raise HTTPException(
            status_code=403,
            detail="Vui lòng cập nhật tài khoản đang đăng nhập qua hồ sơ cá nhân. Bạn không thể tự thay đổi email."
        )

    user = db.query(User).filter(
        User.user_id == user_id
    ).first()

    if user is None:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy người dùng"
        )

    if data.email is not None:
        duplicate_email = db.query(User).filter(
            User.email == data.email,
            User.user_id != user_id
        ).first()

        if duplicate_email:
            raise HTTPException(
                status_code=400,
                detail="Email đã tồn tại trong hệ thống"
            )

        user.email = data.email

    if data.full_name is not None:
        user.full_name = data.full_name

    if data.phone is not None:
        user.phone = data.phone

    db.commit()

    return {
        "message": "Cập nhật tài khoản thành công",
        "user_id": user.user_id
    }


@router.delete("/{user_id}")
def delete_user(
    user_id: int,
    current_user: User = Depends(
        require_permission("USER_MANAGE")
    ),
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(
        User.user_id == user_id
    ).first()

    if user is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Không tìm thấy người dùng"
        )

    if current_user.user_id == user_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Bạn không thể xóa tài khoản hiện tại của chính mình"
        )

    try:
        # Xóa các ràng buộc phụ thuộc
        db.query(UserRole).filter(UserRole.user_id == user_id).delete()
        db.query(UserSession).filter(UserSession.user_id == user_id).delete()
        db.query(PasswordResetToken).filter(PasswordResetToken.user_id == user_id).delete()

        try:
            db.execute(
                text("UPDATE classes SET instructor_id = NULL WHERE instructor_id = :user_id"),
                {"user_id": user_id}
            )
        except Exception:
            pass

        try:
            db.execute(
                text("DELETE FROM student_classes WHERE student_id = :user_id"),
                {"user_id": user_id}
            )
        except Exception:
            pass

        db.delete(user)
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Lỗi khi xóa tài khoản: {str(e)}"
        )

    return {
        "message": "Xóa tài khoản thành công",
        "user_id": user_id
    }


@router.put("/{user_id}/lock")
def lock_user(
    user_id: int,
    data: LockUserRequest,
    current_user: User = Depends(
        require_permission("USER_MANAGE")
    ),
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(
        User.user_id == user_id
    ).first()

    if user is None:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy người dùng"
        )

    if current_user.user_id == user_id:
        raise HTTPException(
            status_code=400,
            detail="Bạn không thể tự khóa tài khoản của chính mình"
        )

    lock_reason = data.lock_reason.strip()

    if not lock_reason:
        raise HTTPException(
            status_code=400,
            detail="Bắt buộc nhập lý do khóa tài khoản"
        )

    if user.is_locked:
        raise HTTPException(
            status_code=400,
            detail="Tài khoản đã bị khóa"
        )

    assigned_classes = db.execute(
        text(
            """
            SELECT
                class_id,
                class_name
            FROM classes
            WHERE instructor_id = :user_id
            ORDER BY class_id ASC
            """
        ),
        {
            "user_id": user_id
        }
    ).mappings().all()

    user.is_locked = True
    user.lock_reason = lock_reason

    db.query(UserSession).filter(
        UserSession.user_id == user_id,
        UserSession.revoked == False
    ).update(
        {
            UserSession.revoked: True
        },
        synchronize_session=False
    )

    db.commit()

    classes_need_handover = [
        {
            "class_id": item["class_id"],
            "class_name": item["class_name"]
        }
        for item in assigned_classes
    ]

    return {
        "message": "Khóa tài khoản thành công",
        "user_id": user.user_id,
        "lock_reason": user.lock_reason,
        "classes_need_handover": classes_need_handover,
        "warning": (
            "Người dùng đang phụ trách lớp học. "
            "Cần thực hiện bàn giao."
            if classes_need_handover
            else None
        )
    }


@router.put("/{user_id}/unlock")
def unlock_user(
    user_id: int,
    current_user: User = Depends(
        require_permission("USER_MANAGE")
    ),
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(
        User.user_id == user_id
    ).first()

    if user is None:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy người dùng"
        )

    if not user.is_locked:
        raise HTTPException(
            status_code=400,
            detail="Tài khoản hiện không bị khóa"
        )

    user.is_locked = False
    user.lock_reason = None
    user.failed_login_attempts = 0
    user.locked_until = None

    db.commit()

    return {
        "message": "Mở khóa tài khoản thành công",
        "user_id": user.user_id
    }


@router.get("")
def get_users(
    search: str | None = Query(
        default=None
    ),
    role_id: int | None = Query(
        default=None
    ),
    status: str | None = Query(
        default=None
    ),
    page: int = Query(
        default=1,
        ge=1
    ),
    page_size: int = Query(
        default=20,
        ge=1,
        le=100
    ),
    current_user: User = Depends(
        require_permission("USER_MANAGE")
    ),
    db: Session = Depends(get_db)
):
    query = db.query(User)

    if search:
        keyword = f"%{search}%"

        query = query.filter(
            or_(
                User.full_name.like(keyword),
                User.email.like(keyword),
                User.phone.like(keyword)
            )
        )

    if role_id is not None:
        query = query.join(
            UserRole,
            User.user_id == UserRole.user_id
        ).filter(
            UserRole.role_id == role_id
        )

    if status == "locked":
        query = query.filter(
            User.is_locked == True
        )

    elif status == "active":
        query = query.filter(
            User.is_locked == False
        )

    total = query.count()

    users = (
        query
        .order_by(User.user_id.asc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    user_ids = [
        user.user_id
        for user in users
    ]

    role_rows = []

    if user_ids:
        role_rows = (
            db.query(
                UserRole.user_id,
                Role.role_id,
                Role.role_name
            )
            .join(
                Role,
                UserRole.role_id == Role.role_id
            )
            .filter(
                UserRole.user_id.in_(user_ids)
            )
            .order_by(
                UserRole.user_id.asc(),
                Role.role_id.asc()
            )
            .all()
        )

    roles_by_user = {}

    for row in role_rows:
        if row.user_id not in roles_by_user:
            roles_by_user[row.user_id] = []

        roles_by_user[row.user_id].append(
            {
                "role_id": row.role_id,
                "role_name": row.role_name
            }
        )

    return {
        "page": page,
        "page_size": page_size,
        "total": total,
        "items": [
            {
                "user_id": user.user_id,
                "username": user.username,
                "full_name": user.full_name,
                "email": user.email,
                "phone": user.phone,
                "role_id": user.role_id,
                "roles": roles_by_user.get(
                    user.user_id,
                    []
                ),
                "status": (
                    "locked"
                    if user.is_locked
                    else "active"
                ),
                "lock_reason": user.lock_reason
            }
            for user in users
        ]
    }


ROLE_NAME_MAP = {
    # ADMIN
    "ADMIN": (1, "Quản trị viên", "ADMIN"),
    "QUẢN TRỊ VIÊN": (1, "Quản trị viên", "ADMIN"),
    "QUAN TRI VIEN": (1, "Quản trị viên", "ADMIN"),
    # INSTRUCTOR
    "INSTRUCTOR": (2, "Giảng viên", "INSTRUCTOR"),
    "GIẢNG VIÊN": (2, "Giảng viên", "INSTRUCTOR"),
    "GIANG VIEN": (2, "Giảng viên", "INSTRUCTOR"),
    "TEACHER": (2, "Giảng viên", "INSTRUCTOR"),
    # STUDENT
    "STUDENT": (3, "Học viên", "STUDENT"),
    "HỌC VIÊN": (3, "Học viên", "STUDENT"),
    "HOC VIEN": (3, "Học viên", "STUDENT"),
    "SINH VIÊN": (3, "Học viên", "STUDENT"),
    "SINH VIEN": (3, "Học viên", "STUDENT"),
    # ACCOUNTANT
    "ACCOUNTANT": (4, "Kế toán", "ACCOUNTANT"),
    "KẾ TOÁN": (4, "Kế toán", "ACCOUNTANT"),
    "KE TOAN": (4, "Kế toán", "ACCOUNTANT"),
    # TRAINING_MANAGER
    "TRAINING_MANAGER": (5, "Quản lý đào tạo", "TRAINING_MANAGER"),
    "MANAGER": (5, "Quản lý đào tạo", "TRAINING_MANAGER"),
    "QUẢN LÝ ĐÀO TẠO": (5, "Quản lý đào tạo", "TRAINING_MANAGER"),
    "QUAN LY DAO TAO": (5, "Quản lý đào tạo", "TRAINING_MANAGER"),
    "QUẢN LÝ": (5, "Quản lý đào tạo", "TRAINING_MANAGER"),
    "QUAN LY": (5, "Quản lý đào tạo", "TRAINING_MANAGER"),
    # ADMISSIONS
    "ADMISSIONS": (6, "Tuyển sinh", "ADMISSIONS"),
    "TUYỂN SINH": (6, "Tuyển sinh", "ADMISSIONS"),
    "TUYEN SINH": (6, "Tuyển sinh", "ADMISSIONS"),
    # ACADEMIC_AFFAIRS
    "ACADEMIC_AFFAIRS": (7, "Giáo vụ", "ACADEMIC_AFFAIRS"),
    "GIÁO VỤ": (7, "Giáo vụ", "ACADEMIC_AFFAIRS"),
    "GIAO VU": (7, "Giáo vụ", "ACADEMIC_AFFAIRS"),
    # MANAGEMENT
    "MANAGEMENT": (8, "Ban quản lý", "MANAGEMENT"),
    "BAN QUẢN LÝ": (8, "Ban quản lý", "MANAGEMENT"),
    "BAN QUAN LY": (8, "Ban quản lý", "MANAGEMENT"),
    "BAN GIÁM HIỆU": (8, "Ban quản lý", "MANAGEMENT"),
    "BAN GIAM HIEU": (8, "Ban quản lý", "MANAGEMENT"),
}


def resolve_role_info(role_raw: str, db: Session):
    clean = str(role_raw).strip()
    if not clean:
        return None

    clean_upper = clean.upper()
    if clean_upper in ROLE_NAME_MAP:
        return ROLE_NAME_MAP[clean_upper]

    if clean.isdigit():
        role_id = int(clean)
        role_db = db.query(Role).filter(Role.role_id == role_id).first()
        if role_db:
            display_name = next(
                (v[1] for k, v in ROLE_NAME_MAP.items() if v[0] == role_db.role_id),
                role_db.role_name
            )
            return (role_db.role_id, display_name, role_db.role_name)

    role_db = db.query(Role).filter(Role.role_name.ilike(clean)).first()
    if role_db:
        display_name = next(
            (v[1] for k, v in ROLE_NAME_MAP.items() if v[0] == role_db.role_id),
            role_db.role_name
        )
        return (role_db.role_id, display_name, role_db.role_name)

    return None


def find_column(df_columns, candidates):
    normalized_cols = {str(col).strip().lower(): col for col in df_columns}
    for c in candidates:
        if c.lower() in normalized_cols:
            return normalized_cols[c.lower()]
    return None


EMAIL_REGEX = re.compile(r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$")


@router.get("/import-excel/template")
def download_import_template(
    current_user: User = Depends(require_permission("USER_MANAGE")),
    db: Session = Depends(get_db)
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MauNhapNguoiDung"

    headers = ["HoTen", "Email", "SoDienThoai", "VaiTro"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    sample_rows = [
        ["Nguyễn Văn An", "an.nguyen@example.com", "0912345678", "Học viên"],
        ["Trần Thị Bình", "binh.tran@example.com", "0987654321", "STUDENT"],
        ["Lê Văn Cường", "cuong.le@example.com", "0901234567", "Giảng viên"],
        ["Phạm Thị Dung", "dung.pham@example.com", "0918765432", "Quản trị viên"],
    ]

    for row_data in sample_rows:
        ws.append(row_data)

    for row in ws.iter_rows(min_row=2, max_row=len(sample_rows) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    column_widths = {"A": 26, "B": 32, "C": 18, "D": 22}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="mau_nhap_nguoi_dung.xlsx"'
        }
    )


def parse_and_validate_excel_rows(df: pd.DataFrame, db: Session):
    col_name = find_column(df.columns, ["HoTen", "Họ và tên", "Họ tên", "FullName", "Full Name", "Tên"])
    col_email = find_column(df.columns, ["Email", "Địa chỉ email", "Thu dien tu", "Mail"])
    col_phone = find_column(df.columns, ["SoDienThoai", "Số điện thoại", "SDT", "SĐT", "Phone"])
    col_role = find_column(df.columns, ["VaiTro", "Vai trò", "Role", "Chức vụ"])

    if not col_name or not col_email or not col_role:
        missing = []
        if not col_name:
            missing.append("HoTen")
        if not col_email:
            missing.append("Email")
        if not col_role:
            missing.append("VaiTro")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Thiếu cột bắt buộc trong tệp Excel: {', '.join(missing)}"
        )

    existing_db_emails = {
        email.lower() for (email,) in db.query(User.email).all() if email
    }

    # Count email occurrences in excel to detect duplicates
    email_counts = {}
    for _, row in df.iterrows():
        raw_e = str(row[col_email]).strip().lower() if pd.notna(row[col_email]) else ""
        if raw_e:
            email_counts[raw_e] = email_counts.get(raw_e, 0) + 1

    parsed_rows = []
    seen_file_emails = set()

    for index, row in df.iterrows():
        line_num = index + 2
        errors = []

        raw_name = str(row[col_name]).strip() if pd.notna(row[col_name]) else ""
        raw_email = str(row[col_email]).strip() if pd.notna(row[col_email]) else ""
        raw_phone = str(row[col_phone]).strip() if col_phone and pd.notna(row[col_phone]) else ""
        raw_role = str(row[col_role]).strip() if pd.notna(row[col_role]) else ""

        # Validate HoTen
        if not raw_name:
            errors.append("Thiếu Họ và tên")
        elif len(raw_name) > 100:
            errors.append("Họ và tên vượt quá 100 ký tự")

        # Validate Email
        if not raw_email:
            errors.append("Thiếu Email")
        else:
            email_lower = raw_email.lower()
            if not EMAIL_REGEX.match(raw_email):
                errors.append("Email không đúng định dạng")
            elif email_lower in existing_db_emails:
                errors.append("Email đã tồn tại trong hệ thống")
            elif email_counts.get(email_lower, 0) > 1:
                if email_lower in seen_file_emails:
                    errors.append("Email bị trùng lặp trong chính tệp Excel")
                else:
                    seen_file_emails.add(email_lower)
                    errors.append("Email bị trùng lặp trong chính tệp Excel")
            else:
                seen_file_emails.add(email_lower)

        # Validate SoDienThoai
        clean_phone = raw_phone
        if raw_phone:
            clean_digits = re.sub(r"[\s.-]", "", raw_phone)
            if clean_digits.endswith(".0"):
                clean_digits = clean_digits[:-2]
            if len(clean_digits) == 9 and clean_digits.isdigit() and not clean_digits.startswith(("0", "+84")):
                clean_digits = "0" + clean_digits
            if not re.match(r"^(0|\+84)[0-9]{8,11}$", clean_digits):
                errors.append("Số điện thoại không hợp lệ")
            else:
                clean_phone = clean_digits

        # Validate VaiTro
        role_info = None
        if not raw_role:
            errors.append("Thiếu Vai trò")
        else:
            role_info = resolve_role_info(raw_role, db)
            if not role_info:
                errors.append(
                    "Vai trò không hợp lệ (hỗ trợ: Học viên, Giảng viên, Quản trị viên, Quản lý đào tạo, Kế toán, Tuyển sinh, Giáo vụ, Ban quản lý)"
                )

        role_id = role_info[0] if role_info else None
        role_name_display = role_info[1] if role_info else (raw_role or "Chưa xác định")

        is_valid = len(errors) == 0

        parsed_rows.append(
            ImportUserRow(
                row_index=line_num,
                full_name=raw_name,
                email=raw_email,
                phone=clean_phone if clean_phone else None,
                role_input=raw_role,
                role_name=role_name_display,
                role_id=role_id,
                is_valid=is_valid,
                errors=errors
            )
        )

    valid_count = sum(1 for r in parsed_rows if r.is_valid)
    invalid_count = len(parsed_rows) - valid_count

    return parsed_rows, valid_count, invalid_count


@router.post("/import-excel/preview", response_model=ImportPreviewResponse)
async def preview_users_from_excel(
    file: UploadFile = File(...),
    current_user: User = Depends(require_permission("USER_MANAGE")),
    db: Session = Depends(get_db)
):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Tệp không đúng định dạng Excel (.xlsx, .xls)"
        )

    contents = await file.read()
    try:
        df = pd.read_excel(io.BytesIO(contents), dtype=str)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không thể đọc nội dung tệp Excel. Vui lòng kiểm tra định dạng tệp."
        )

    if df.empty:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Tệp Excel không chứa dữ liệu dòng nào"
        )

    parsed_rows, valid_count, invalid_count = parse_and_validate_excel_rows(df, db)

    return ImportPreviewResponse(
        total_rows=len(parsed_rows),
        valid_count=valid_count,
        invalid_count=invalid_count,
        rows=parsed_rows
    )


@router.post("/import-excel/confirm", response_model=ImportSummaryResponse)
def confirm_import_users(
    payload: ConfirmImportRequest,
    current_user: User = Depends(require_permission("USER_MANAGE")),
    db: Session = Depends(get_db)
):
    success_count = 0
    failed_count = 0
    errors = []
    default_password_hash = hash_password("12345678")

    for item in payload.users:
        # Kiểm tra trùng email trong DB tại thời điểm nhập
        existing_email = db.query(User).filter(User.email == item.email).first()
        if existing_email:
            failed_count += 1
            errors.append(f"Dòng {item.row_index}: Email '{item.email}' đã tồn tại trong hệ thống")
            continue

        # Kiểm tra vai trò
        role = db.query(Role).filter(Role.role_id == item.role_id).first()
        if not role:
            failed_count += 1
            errors.append(f"Dòng {item.row_index}: Vai trò ID {item.role_id} không tồn tại")
            continue

        # Tự sinh username an toàn từ email
        base_username = item.email.split("@")[0].strip().lower()
        base_username = re.sub(r"[^a-zA-Z0-9_]", "_", base_username) or "user"
        username = base_username
        suffix = 1
        while db.query(User).filter(User.username == username).first():
            username = f"{base_username}_{suffix}"
            suffix += 1

        try:
            new_user = User(
                username=username,
                full_name=item.full_name,
                email=item.email,
                phone=item.phone if item.phone else None,
                password=default_password_hash,
                role_id=item.role_id,
                is_locked=False,
                failed_login_attempts=0
            )
            db.add(new_user)
            db.flush()

            user_role = UserRole(
                user_id=new_user.user_id,
                role_id=item.role_id
            )
            db.add(user_role)
            db.commit()
            success_count += 1
        except Exception as e:
            db.rollback()
            failed_count += 1
            errors.append(f"Dòng {item.row_index}: Lỗi khi lưu dữ liệu ({str(e)})")

    total_rows = len(payload.users)
    summary_text = f"Tổng số: {total_rows} | Thành công: {success_count} | Bỏ qua: {failed_count}"

    return ImportSummaryResponse(
        message="Nhập người dùng hoàn tất",
        total_rows=total_rows,
        success_count=success_count,
        failed_count=failed_count,
        errors=errors,
        summary_text=summary_text
    )


@router.post("/import-excel", response_model=ImportSummaryResponse)
async def import_users_from_excel(
    file: UploadFile = File(...),
    current_user: User = Depends(require_permission("USER_MANAGE")),
    db: Session = Depends(get_db)
):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Tệp không đúng định dạng Excel (.xlsx, .xls)"
        )

    contents = await file.read()
    try:
        df = pd.read_excel(io.BytesIO(contents), dtype=str)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không thể đọc nội dung tệp Excel"
        )

    if df.empty:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Tệp Excel không chứa dữ liệu dòng nào"
        )

    parsed_rows, _, _ = parse_and_validate_excel_rows(df, db)

    valid_items = [
        ImportConfirmUserItem(
            row_index=r.row_index,
            full_name=r.full_name,
            email=r.email,
            phone=r.phone,
            role_id=r.role_id  # type: ignore
        )
        for r in parsed_rows
        if r.is_valid and r.role_id is not None
    ]

    invalid_row_errors = [
        f"Dòng {r.row_index} ({r.email or 'Không có email'}): {'; '.join(r.errors)}"
        for r in parsed_rows
        if not r.is_valid
    ]

    confirm_res = confirm_import_users(
        payload=ConfirmImportRequest(users=valid_items),
        current_user=current_user,
        db=db
    )

    all_errors = invalid_row_errors + confirm_res.errors
    total_rows = len(parsed_rows)
    failed_count = len(parsed_rows) - confirm_res.success_count

    return ImportSummaryResponse(
        message="Nhập dữ liệu hoàn tất",
        total_rows=total_rows,
        success_count=confirm_res.success_count,
        failed_count=failed_count,
        errors=all_errors,
        summary_text=f"Tổng số: {total_rows} | Thành công: {confirm_res.success_count} | Bỏ qua: {failed_count}"
    )
