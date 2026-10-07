import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import app
from app.database import get_db, SessionLocal
from app.core.security import create_access_token, hash_password
from app.models.user import User
from app.models.role import Role
from app.models.user_role import UserRole


client = TestClient(app)


def create_excel_bytes(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@pytest.fixture(scope="module")
def admin_token():
    db: Session = SessionLocal()
    # Check if admin user exists, or create one for test
    admin_user = db.query(User).filter(User.username == "admin_test_scrum29").first()
    if not admin_user:
        admin_user = User(
            username="admin_test_scrum29",
            email="admin_test_scrum29@test.com",
            full_name="Admin Test SCRUM-29",
            password=hash_password("123456"),
            role_id=1,
            is_locked=False,
            failed_login_attempts=0
        )
        db.add(admin_user)
        db.flush()
        db.add(UserRole(user_id=admin_user.user_id, role_id=1))
        db.commit()
        db.refresh(admin_user)

    token = create_access_token(
        user_id=admin_user.user_id,
        role="ADMIN"
    )
    yield token

    # Cleanup admin test user
    db.query(UserRole).filter(UserRole.user_id == admin_user.user_id).delete()
    db.query(User).filter(User.user_id == admin_user.user_id).delete()
    db.commit()
    db.close()


@pytest.fixture(scope="module")
def student_token():
    db: Session = SessionLocal()
    student_user = db.query(User).filter(User.username == "student_test_scrum29").first()
    if not student_user:
        student_user = User(
            username="student_test_scrum29",
            email="student_test_scrum29@test.com",
            full_name="Student Test SCRUM-29",
            password=hash_password("123456"),
            role_id=3,
            is_locked=False,
            failed_login_attempts=0
        )
        db.add(student_user)
        db.flush()
        db.add(UserRole(user_id=student_user.user_id, role_id=3))
        db.commit()
        db.refresh(student_user)

    token = create_access_token(
        user_id=student_user.user_id,
        role="STUDENT"
    )
    yield token

    db.query(UserRole).filter(UserRole.user_id == student_user.user_id).delete()
    db.query(User).filter(User.user_id == student_user.user_id).delete()
    db.commit()
    db.close()


def test_download_template_unauthorized():
    response = client.get("/api/users/import-excel/template")
    assert response.status_code in [401, 403]


def test_download_template_forbidden_for_student(student_token):
    response = client.get(
        "/api/users/import-excel/template",
        headers={"Authorization": f"Bearer {student_token}"}
    )
    assert response.status_code == 403


def test_download_template_success(admin_token):
    response = client.get(
        "/api/users/import-excel/template",
        headers={"Authorization": f"Bearer {admin_token}"}
    )
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
    
    # Read response as Excel
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    first_row = [cell.value for cell in ws[1]]
    assert first_row == ["HoTen", "Email", "SoDienThoai", "VaiTro"]


def test_preview_invalid_file_format(admin_token):
    response = client.post(
        "/api/users/import-excel/preview",
        headers={"Authorization": f"Bearer {admin_token}"},
        files={"file": ("test.txt", b"plain text", "text/plain")}
    )
    assert response.status_code == 400
    assert "định dạng" in response.json()["detail"].lower()


def test_preview_missing_required_columns(admin_token):
    excel_data = create_excel_bytes(["HoTen", "SoDienThoai"], [["Nguyễn Văn A", "0912345678"]])
    response = client.post(
        "/api/users/import-excel/preview",
        headers={"Authorization": f"Bearer {admin_token}"},
        files={"file": ("test.xlsx", excel_data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 400
    assert "Thiếu cột bắt buộc" in response.json()["detail"]


def test_preview_validations(admin_token):
    headers = ["HoTen", "Email", "SoDienThoai", "VaiTro"]
    rows = [
        # Dòng 2: Hợp lệ (Học viên)
        ["Lê Hoàng Yến", "lehoangyen_valid@example.com", "0911223344", "Học viên"],
        # Dòng 3: Thiếu họ tên
        ["", "noname@example.com", "0922334455", "STUDENT"],
        # Dòng 4: Email sai định dạng
        ["Trần Văn Sai", "invalid-email-address", "0933445566", "Học viên"],
        # Dòng 5: Vai trò không hợp lệ
        ["Phạm Văn Role", "invalidrole@example.com", "0944556677", "SUPER_ADMIN_CUSTOM"],
        # Dòng 6: Email trùng lặp trong chính Excel (lần 1)
        ["Vũ Trùng A", "duplicate_excel@example.com", "0955667788", "Giảng viên"],
        # Dòng 7: Email trùng lặp trong chính Excel (lần 2)
        ["Vũ Trùng B", "duplicate_excel@example.com", "0955667789", "Giảng viên"],
    ]

    excel_data = create_excel_bytes(headers, rows)
    response = client.post(
        "/api/users/import-excel/preview",
        headers={"Authorization": f"Bearer {admin_token}"},
        files={"file": ("users.xlsx", excel_data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )

    assert response.status_code == 200
    data = response.json()

    assert data["total_rows"] == 6
    assert data["valid_count"] == 1
    assert data["invalid_count"] == 5

    preview_rows = data["rows"]
    # Row 2 is valid
    assert preview_rows[0]["is_valid"] is True
    assert preview_rows[0]["role_id"] == 3
    assert preview_rows[0]["role_name"] == "Học viên"

    # Row 3 lacks name
    assert preview_rows[1]["is_valid"] is False
    assert any("Họ và tên" in err for err in preview_rows[1]["errors"])

    # Row 4 invalid email format
    assert preview_rows[2]["is_valid"] is False
    assert any("định dạng" in err for err in preview_rows[2]["errors"])

    # Row 5 invalid role
    assert preview_rows[3]["is_valid"] is False
    assert any("Vai trò không hợp lệ" in err for err in preview_rows[3]["errors"])

    # Row 6 & 7 duplicate email in Excel
    assert preview_rows[4]["is_valid"] is False
    assert any("trùng lặp trong chính tệp Excel" in err for err in preview_rows[4]["errors"])
    assert preview_rows[5]["is_valid"] is False
    assert any("trùng lặp trong chính tệp Excel" in err for err in preview_rows[5]["errors"])


def test_confirm_import_isolated_and_independent(admin_token):
    db: Session = SessionLocal()

    # Pre-existing user to trigger duplicate email in DB
    existing_email = "pre_existing_user@example.com"
    existing_user = db.query(User).filter(User.email == existing_email).first()
    if not existing_user:
        existing_user = User(
            username="pre_existing_user",
            email=existing_email,
            full_name="Người Dùng Đã Có",
            password=hash_password("123456"),
            role_id=3,
            is_locked=False,
            failed_login_attempts=0
        )
        db.add(existing_user)
        db.commit()

    import_users_payload = {
        "users": [
            {
                "row_index": 2,
                "full_name": "Nguyễn Thành Công 1",
                "email": "thanhcong1_scrum29@example.com",
                "phone": "0911223344",
                "role_id": 3
            },
            {
                "row_index": 3,
                "full_name": "Nguyễn Thất Bại (Trùng DB)",
                "email": existing_email,
                "phone": "0922334455",
                "role_id": 3
            },
            {
                "row_index": 4,
                "full_name": "Nguyễn Thành Công 2",
                "email": "thanhcong2_scrum29@example.com",
                "phone": "0933445566",
                "role_id": 2
            }
        ]
    }

    response = client.post(
        "/api/users/import-excel/confirm",
        headers={"Authorization": f"Bearer {admin_token}"},
        json=import_users_payload
    )

    assert response.status_code == 200
    res_data = response.json()

    assert res_data["total_rows"] == 3
    assert res_data["success_count"] == 2
    assert res_data["failed_count"] == 1
    assert "Tổng số: 3 | Thành công: 2 | Bỏ qua: 1" in res_data["summary_text"]
    assert any(existing_email in err for err in res_data["errors"])

    # Verify users were really created in DB
    u1 = db.query(User).filter(User.email == "thanhcong1_scrum29@example.com").first()
    u2 = db.query(User).filter(User.email == "thanhcong2_scrum29@example.com").first()
    assert u1 is not None
    assert u1.role_id == 3
    assert u2 is not None
    assert u2.role_id == 2

    # Clean up test users
    for u in [u1, u2]:
        if u:
            db.query(UserRole).filter(UserRole.user_id == u.user_id).delete()
            db.query(User).filter(User.user_id == u.user_id).delete()
    db.query(User).filter(User.email == existing_email).delete()
    db.commit()
    db.close()


def test_direct_import_excel_endpoint(admin_token):
    db: Session = SessionLocal()
    headers = ["HoTen", "Email", "SoDienThoai", "VaiTro"]
    direct_email = "direct_import_scrum29@example.com"
    rows = [
        ["Trần Văn Direct", direct_email, "0988776655", "Giảng viên"],
    ]

    excel_data = create_excel_bytes(headers, rows)
    response = client.post(
        "/api/users/import-excel",
        headers={"Authorization": f"Bearer {admin_token}"},
        files={"file": ("direct_users.xlsx", excel_data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )

    assert response.status_code == 200
    res_data = response.json()
    assert res_data["total_rows"] == 1
    assert res_data["success_count"] == 1
    assert res_data["failed_count"] == 0

    imported_user = db.query(User).filter(User.email == direct_email).first()
    assert imported_user is not None
    assert imported_user.role_id == 2
    assert imported_user.full_name == "Trần Văn Direct"

    # Cleanup
    db.query(UserRole).filter(UserRole.user_id == imported_user.user_id).delete()
    db.query(User).filter(User.user_id == imported_user.user_id).delete()
    db.commit()
    db.close()

